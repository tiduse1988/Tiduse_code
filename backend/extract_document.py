#!/usr/bin/env python3
import json
import os
import re
import subprocess
import sys
import tempfile
import base64
import urllib.parse
import urllib.request
from pathlib import Path


def clean_text(value):
    return re.sub(r"[ \t]+", " ", (value or "").replace("\x00", "")).strip()


BAIDU_OCR_TOKEN_CACHE = None


def baidu_ocr_configured():
    return bool(os.environ.get("BAIDU_OCR_API_KEY") and os.environ.get("BAIDU_OCR_SECRET_KEY"))


def baidu_access_token():
    global BAIDU_OCR_TOKEN_CACHE
    if BAIDU_OCR_TOKEN_CACHE:
        return BAIDU_OCR_TOKEN_CACHE
    api_key = os.environ.get("BAIDU_OCR_API_KEY")
    secret_key = os.environ.get("BAIDU_OCR_SECRET_KEY")
    if not api_key or not secret_key:
        raise RuntimeError("未配置百度云OCR API Key或Secret Key")
    query = urllib.parse.urlencode({
        "grant_type": "client_credentials",
        "client_id": api_key,
        "client_secret": secret_key,
    })
    url = f"https://aip.baidubce.com/oauth/2.0/token?{query}"
    with urllib.request.urlopen(url, timeout=20) as response:
        payload = json.loads(response.read().decode("utf-8"))
    token = payload.get("access_token")
    if not token:
        raise RuntimeError(payload.get("error_description") or payload.get("error") or "获取百度云OCR token失败")
    BAIDU_OCR_TOKEN_CACHE = token
    return token


def baidu_ocr_image_bytes(image_bytes):
    token = baidu_access_token()
    endpoint = os.environ.get("BAIDU_OCR_ENDPOINT") or "https://aip.baidubce.com/rest/2.0/ocr/v1/accurate_basic"
    url = f"{endpoint}?access_token={urllib.parse.quote(token)}"
    data = urllib.parse.urlencode({
        "image": base64.b64encode(image_bytes).decode("ascii"),
        "language_type": "CHN_ENG",
        "detect_direction": "true",
        "paragraph": "true",
    }).encode("utf-8")
    request = urllib.request.Request(url, data=data, headers={"Content-Type": "application/x-www-form-urlencoded"})
    with urllib.request.urlopen(request, timeout=45) as response:
        payload = json.loads(response.read().decode("utf-8"))
    if payload.get("error_code"):
        raise RuntimeError(f"百度云OCR失败：{payload.get('error_msg') or payload.get('error_code')}")
    return "\n".join(clean_text(item.get("words")) for item in payload.get("words_result", []) if item.get("words"))


def extract_image_with_baidu(path):
    text = clean_text(baidu_ocr_image_bytes(Path(path).read_bytes()))
    warnings = []
    if not text:
        warnings.append("百度云OCR未识别到有效文字。")
    return [{"page": 1, "text": text, "charCount": len(text), "tableCount": 0, "ocrProvider": "Baidu OCR"}], [], warnings


def render_pdf_pages(path, max_pages):
    with tempfile.TemporaryDirectory() as tmp:
        prefix = os.path.join(tmp, "page")
        result = subprocess.run(
            ["pdftoppm", "-png", "-r", "180", "-f", "1", "-l", str(max_pages), path, prefix],
            capture_output=True,
            text=True,
        )
        if result.returncode != 0:
            raise RuntimeError(result.stderr.strip() or "pdftoppm渲染扫描PDF失败")
        rendered = sorted(Path(tmp).glob("page-*.png"))
        return [(index + 1, item.read_bytes()) for index, item in enumerate(rendered)]


def ocr_pdf_pages_with_baidu(path, existing_pages):
    warnings = []
    if not baidu_ocr_configured():
        return existing_pages, ["未配置百度云OCR，扫描PDF仅完成本地文本提取。"]
    try:
        max_pages = int(os.environ.get("BAIDU_OCR_MAX_PAGES") or "80")
    except ValueError:
        max_pages = 80
    try:
        rendered_pages = render_pdf_pages(path, min(max_pages, max(1, len(existing_pages))))
    except Exception as exc:
        return existing_pages, [f"百度云OCR前置渲染失败：{exc}"]

    page_map = {page.get("page"): dict(page) for page in existing_pages}
    for page_no, image_bytes in rendered_pages:
        current = page_map.get(page_no, {"page": page_no, "text": "", "charCount": 0, "tableCount": 0})
        if current.get("charCount", 0) >= 80:
            continue
        try:
            text = clean_text(baidu_ocr_image_bytes(image_bytes))
            if text:
                current["text"] = text
                current["charCount"] = len(text)
                current["ocrProvider"] = "Baidu OCR"
                page_map[page_no] = current
        except Exception as exc:
            warnings.append(f"第{page_no}页百度云OCR失败：{exc}")
    if len(existing_pages) > max_pages:
        warnings.append(f"百度云OCR已按配置仅处理前{max_pages}页；如需更多页请调整BAIDU_OCR_MAX_PAGES。")
    return [page_map.get(index, page) for index, page in enumerate(existing_pages, 1)], warnings


def detect_sections(full_text):
    patterns = [
        r"招标公告",
        r"投标人须知",
        r"资格条件|资格要求|资格审查",
        r"商务要求|商务条款",
        r"技术要求|技术参数|服务要求",
        r"评分标准|评分办法|评标办法",
        r"废标|无效投标|否决投标",
        r"响应文件格式|投标文件格式|文件格式",
        r"资料清单|附件|证明材料",
        r"采购需求|项目需求|发包人要求|功能要求|工程规模及内容|工程范围|包括的工作|建设规模|建设内容|招标范围|承包范围|服务内容|服务范围|采购内容|采购范围|供货要求|设备清单|清单及参数",
    ]
    sections = []
    for pattern in patterns:
        match = re.search(pattern, full_text)
        if match:
            start = max(0, match.start() - 120)
            end = min(len(full_text), match.end() + 500)
            sections.append({
                "title": match.group(0),
                "offset": match.start(),
                "snippet": clean_text(full_text[start:end])[:650],
            })
    return sections


def extract_pdf(path):
    import pdfplumber

    pages = []
    tables = []
    warnings = []
    with pdfplumber.open(path) as pdf:
        for index, page in enumerate(pdf.pages, 1):
            text = clean_text(page.extract_text() or "")
            page_tables = page.extract_tables() or []
            table_count = 0
            for table in page_tables:
                rows = []
                for row in table:
                    rows.append([clean_text(cell) for cell in row])
                if any(any(cell for cell in row) for row in rows):
                    table_count += 1
                    tables.append({"page": index, "rows": rows[:80]})
            if len(text) < 30:
                warnings.append(f"第{index}页文本较少，可能是扫描页或图片页，建议接入OCR复核。")
            pages.append({"page": index, "text": text, "charCount": len(text), "tableCount": table_count})
    low_text_ratio = sum(1 for page in pages if page.get("charCount", 0) < 30) / max(1, len(pages)) if pages else 0
    if low_text_ratio > 0.2:
        pages, ocr_warnings = ocr_pdf_pages_with_baidu(path, pages)
        warnings.extend(ocr_warnings)
    return pages, tables, warnings


def extract_docx(path):
    from docx import Document

    doc = Document(path)
    blocks = []
    tables = []
    for para in doc.paragraphs:
        text = clean_text(para.text)
        if text:
            style = getattr(para.style, "name", "")
            blocks.append({"style": style, "text": text})
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append([clean_text(cell.text) for cell in row.cells])
        if any(any(cell for cell in row) for row in rows):
            tables.append({"page": None, "rows": rows[:120]})
    text = "\n".join(block["text"] for block in blocks)
    return [{"page": 1, "text": text, "charCount": len(text), "tableCount": len(tables)}], tables, []


def extract_doc(path):
    warnings = []
    if sys.platform != "darwin":
        return [], [], ["DOC老格式当前仅支持在macOS通过textutil转换；建议先转为DOCX。"]
    with tempfile.TemporaryDirectory() as tmp:
        result = subprocess.run(
            ["textutil", "-convert", "txt", "-output", os.path.join(tmp, "out.txt"), path],
            capture_output=True,
            text=True,
        )
        if result.returncode != 0:
            return [], [], [f"DOC转换失败：{result.stderr.strip() or '未知错误'}"]
        text = Path(tmp, "out.txt").read_text("utf-8", errors="ignore")
    text = clean_text(text)
    if not text:
        warnings.append("DOC转换后未提取到有效文本，建议人工复核或转换为DOCX。")
    return [{"page": 1, "text": text, "charCount": len(text), "tableCount": 0}], [], warnings


def main():
    if len(sys.argv) < 2:
        print(json.dumps({"error": "missing file path"}, ensure_ascii=False))
        return 2

    path = sys.argv[1]
    ext = Path(path).suffix.lower()
    warnings = []
    try:
        if ext == ".pdf":
            pages, tables, warnings = extract_pdf(path)
        elif ext == ".docx":
            pages, tables, warnings = extract_docx(path)
        elif ext == ".doc":
            pages, tables, warnings = extract_doc(path)
        elif ext in [".txt", ".md"]:
            text = Path(path).read_text("utf-8", errors="ignore")
            pages, tables = [{"page": 1, "text": clean_text(text), "charCount": len(text), "tableCount": 0}], []
        elif ext in [".png", ".jpg", ".jpeg"]:
            if baidu_ocr_configured():
                pages, tables, warnings = extract_image_with_baidu(path)
            else:
                pages, tables = [], []
                warnings = ["图片格式需要OCR服务；当前环境未配置百度云OCR。"]
        else:
            pages, tables = [], []
            warnings = [f"暂不支持的文件格式：{ext}"]
    except Exception as exc:
        print(json.dumps({"error": str(exc), "pages": [], "tables": [], "warnings": [str(exc)]}, ensure_ascii=False))
        return 1

    full_text = "\n".join(page["text"] for page in pages if page.get("text"))
    result = {
        "fileName": Path(path).name,
        "fileType": ext.lstrip("."),
        "pageCount": len(pages),
        "charCount": len(full_text),
        "tableCount": len(tables),
        "pages": pages[:400],
        "tables": tables[:120],
        "sections": detect_sections(full_text),
        "warnings": warnings,
        "fullText": full_text[:240000],
        "quality": {
            "hasText": len(full_text) > 0,
            "likelyScanned": bool(pages) and sum(1 for page in pages if page.get("charCount", 0) < 30) / max(1, len(pages)) > 0.5,
            "needsOcr": bool(warnings and any("OCR" in item or "扫描" in item for item in warnings)),
            "textCoverage": round(sum(1 for page in pages if page.get("charCount", 0) >= 30) / max(1, len(pages)), 2) if pages else 0,
        },
    }
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
